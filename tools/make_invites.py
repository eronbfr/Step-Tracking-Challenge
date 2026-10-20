"""Builds invites.html: one-tap personal WhatsApp invites for each teammate.
Reads names + phone numbers from the local spreadsheet. The output contains phone numbers — it is git-ignored.
Run: python tools/make_invites.py [path/to/step-tracking_Team8.xlsx]
"""
import glob
import html
import os
import re
import sys
import unicodedata
import urllib.parse

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
APP_URL = 'https://eronbfr.github.io/Step-Tracking-Challenge/'


def first_name_id(name):
    first = name.split()[0]
    return unicodedata.normalize('NFKD', first).encode('ascii', 'ignore').decode().lower()


def message(first):
    return (f"Hi {first}! 👟 The Team 8 Step Challenge starts today (Oct 20 – Nov 15, 10K steps a day).\n\n"
            f"📲 Install our app: {APP_URL}\n"
            "1️⃣ Pick your name and create a 4-digit PIN\n"
            "2️⃣ Turn on Auto-sync (⚡ tab) so your steps update by themselves\n\n"
            "Let's go, Team 8! 🚀")


def main():
    path = sys.argv[1] if len(sys.argv) > 1 else (glob.glob(os.path.join(ROOT, 'step-tracking_Team8*.xlsx')) or [''])[0]
    if not os.path.exists(path):
        sys.exit('Spreadsheet not found. Pass its path as an argument.')
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    people = []
    for row in range(3, ws.max_row + 1):
        name, phone = ws.cell(row, 2).value, ws.cell(row, 3).value
        if not name or not phone or 'total' in str(name).lower():
            continue
        digits = re.sub(r'\D', '', str(phone))
        first = str(name).split()[0]
        url = f"https://wa.me/{digits}?text={urllib.parse.quote(message(first))}"
        people.append((str(name).strip(), first_name_id(str(name)), url))

    cards = '\n'.join(
        f'<a class="card" href="{html.escape(u)}" target="_blank" rel="noopener"><b>{html.escape(n)}</b><span>Send WhatsApp invite →</span></a>'
        for n, _, u in people)
    page = f"""<!doctype html><html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1"><title>Team 8 invites</title>
<style>
body{{font:15px/1.4 system-ui,sans-serif;background:#f3f5fa;color:#0e1222;margin:0;padding:24px 16px}}
main{{max-width:520px;margin:auto;display:grid;gap:10px}} h1{{margin:0 0 4px}} p{{color:#4a5068;margin:0 0 12px}}
.card{{display:flex;justify-content:space-between;align-items:center;gap:12px;padding:14px 16px;border-radius:16px;
background:#fff;text-decoration:none;color:inherit;box-shadow:0 4px 16px rgba(16,24,48,.08)}}
.card span{{color:#128c7e;font-weight:700;white-space:nowrap}}
</style></head><body><main><h1>👟 Team 8 invites</h1>
<p>One tap opens WhatsApp with a personal message. This file has phone numbers — keep it private.</p>
{cards}</main></body></html>"""
    out = os.path.join(ROOT, 'invites.html')
    with open(out, 'w', encoding='utf-8') as f:
        f.write(page)
    print(f'{len(people)} invites written to {out}')


if __name__ == '__main__':
    main()
